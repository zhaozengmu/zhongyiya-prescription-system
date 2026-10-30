"""历史查询与统计视图（含驾驶舱首页）。"""
from datetime import timedelta

from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.core.paginator import Paginator
from django.db.models import Count, F, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.http import require_GET

from apps.prescriptions.models import Herb, HerbBatch, Prescription, PrescriptionItem
from apps.reviews.models import ReviewRecord
from apps.users.permissions import deny, get_client_ip, has_permission

CACHE_KEY = "dashboard_metrics"
CACHE_TIMEOUT = 120  # 秒，驾驶舱指标允许 2 分钟延迟


def _dashboard_metrics():
    """计算驾驶舱指标（带缓存，避免每次回首页都跑十几次聚合）。"""
    cached = cache.get(CACHE_KEY)
    if cached:
        return cached

    today = timezone.now().date()
    yesterday = today - timedelta(days=1)

    visible = Prescription.objects.filter(is_deleted=False)
    today_presc = visible.filter(created_at__date=today).count()
    yesterday_presc = visible.filter(created_at__date=yesterday).count()
    delta = 0
    if yesterday_presc:
        delta = round((today_presc - yesterday_presc) / yesterday_presc * 100)
    elif today_presc:
        delta = 100

    trend = [
        {
            "date": (today - timedelta(days=i)).strftime("%m-%d"),
            "count": visible.filter(created_at__date=today - timedelta(days=i)).count(),
        }
        for i in range(6, -1, -1)
    ]

    total = visible.count()
    approved = visible.filter(status__in=["APPROVED", "DISPENSED"]).count()
    pass_rate = round(approved / total * 100) if total else 0

    metrics = {
        "today_presc": today_presc,
        "day_delta": delta,
        "pending_review": visible.filter(
            status__in=["SUBMITTED", "PRECHECKED", "ALERTED", "REVIEWING", "FINAL_REVIEWING"]
        ).count(),
        "rejected": visible.filter(status="REJECTED").count(),
        "total_presc": total,
        "pass_rate": pass_rate,
        "unpaid": visible.filter(status="APPROVED", pay_status="UNPAID").count(),
        "to_dispense": visible.filter(status="APPROVED", pay_status="PAID").count(),
        "trend": trend,
        "doctor_rank": list(
            visible.values("doctor__user__first_name", "doctor__user__username")
            .annotate(cnt=Count("id"))
            .order_by("-cnt")[:10]
        ),
        "herb_top": list(
            PrescriptionItem.objects.filter(prescription__is_deleted=False)
            .values("herb__name")
            .annotate(total=Sum("dose"))
            .order_by("-total")[:10]
        ),
        "reject_reasons": list(
            ReviewRecord.objects.exclude(reject_reason="")
            .values("reject_reason")
            .annotate(cnt=Count("id"))
            .order_by("-cnt")[:8]
        ),
        "status_flow": list(
            visible.values("status").annotate(cnt=Count("id")).order_by("-cnt")
        ),
    }
    cache.set(CACHE_KEY, metrics, CACHE_TIMEOUT)
    return metrics


def invalidate_dashboard_cache():
    cache.delete(CACHE_KEY)


@login_required
def dashboard(request):
    """管理驾驶舱首页（需要统计权限，不再对患者开放）。"""
    if not has_permission(request.user, "统计:读"):
        return deny(request, perm_code="统计:读")

    metrics = _dashboard_metrics()
    today = timezone.now().date()
    labels = dict(Prescription.STATUS_CHOICES)

    near_expiry = HerbBatch.objects.filter(
        expire_at__lte=today + timedelta(days=90), expire_at__gte=today
    ).select_related("herb", "supplier")[:10]
    stock_shortage = Herb.objects.filter(stock__lte=F("stock_warning"), stock_warning__gt=0)[:10]

    context = {
        **metrics,
        "near_expiry": near_expiry,
        "stock_shortage": stock_shortage,
        "status_flow": [
            {"label": labels.get(s["status"], s["status"]), "status": s["status"], "cnt": s["cnt"]}
            for s in metrics["status_flow"]
        ],
    }
    return render(request, "stats/dashboard.html", context)


@login_required
def search_view(request):
    """历史处方多条件检索页面。"""
    if not has_permission(request.user, "统计:读"):
        return deny(request, perm_code="统计:读")
    return render(request, "stats/search.html", {"status_choices": Prescription.STATUS_CHOICES})


@login_required
def logs_view(request):
    """操作日志查询页面（分页）。"""
    if not has_permission(request.user, "日志:读"):
        return deny(request, perm_code="日志:读")
    from .models import OperationLog

    qs = OperationLog.objects.select_related("operator").order_by("-created_at")
    operator = request.GET.get("operator", "").strip()
    action = request.GET.get("action", "").strip()
    if operator:
        qs = qs.filter(operator__username__icontains=operator)
    if action:
        qs = qs.filter(action_type__icontains=action)

    page = Paginator(qs, 50).get_page(request.GET.get("page", 1))
    return render(request, "stats/logs.html", {
        "page": page, "logs": page.object_list, "operator": operator, "action": action,
    })


# ============ API 接口 ============

@require_GET
def api_stats_overview(request):
    """GET /api/v1/stats/overview 驾驶舱核心指标。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "统计:读"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    metrics = _dashboard_metrics()
    today = timezone.now().date()
    return JsonResponse({
        "code": 200,
        "message": "ok",
        "data": {
            **metrics,
            "near_expiry": HerbBatch.objects.filter(
                expire_at__lte=today + timedelta(days=90), expire_at__gte=today
            ).count(),
            "stock_shortage": Herb.objects.filter(stock__lte=F("stock_warning"), stock_warning__gt=0).count(),
        },
    })


@require_GET
def api_search(request):
    """GET /api/v1/search/prescriptions 多条件检索（分页）。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "统计:读"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    qs = Prescription.objects.filter(is_deleted=False).select_related("patient", "doctor__user")

    status = request.GET.get("status", "")
    doctor_name = request.GET.get("doctor", "")
    patient_name = request.GET.get("patient", "")
    start = request.GET.get("start", "")
    end = request.GET.get("end", "")
    hit_rule = request.GET.get("hit_rule", "")

    if status:
        qs = qs.filter(status=status)
    if doctor_name:
        qs = qs.filter(doctor__user__username__icontains=doctor_name)
    if patient_name:
        qs = qs.filter(patient__name__icontains=patient_name)
    if start:
        qs = qs.filter(created_at__date__gte=start)
    if end:
        qs = qs.filter(created_at__date__lte=end)
    if hit_rule:
        qs = qs.filter(review_records__precheck_result__icontains=hit_rule).distinct()

    page_no = _safe_int(request.GET.get("page"), 1, 1, 10000)
    page_size = _safe_int(request.GET.get("page_size"), 20, 1, 100)
    page = Paginator(qs.order_by("-created_at"), page_size).get_page(page_no)

    items = [
        {
            "id": p.id,
            "presc_no": p.presc_no,
            "patient_name": p.patient.name,
            "doctor_name": p.doctor.user.get_full_name() or p.doctor.user.username,
            "status": p.get_status_display(),
            "pay_status": p.get_pay_status_display(),
            "created_at": p.created_at.strftime("%Y-%m-%d %H:%M"),
        }
        for p in page.object_list
    ]
    return JsonResponse({
        "code": 200, "message": "ok",
        "data": {"items": items, "total": page.paginator.count, "page": page_no, "page_size": page_size},
    })


@require_GET
def api_logs(request):
    """GET /api/v1/logs 操作日志查询（分页）。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "日志:读"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    from .models import OperationLog

    qs = OperationLog.objects.select_related("operator").order_by("-created_at")
    operator = request.GET.get("operator", "")
    action = request.GET.get("action", "")
    if operator:
        qs = qs.filter(operator__username__icontains=operator)
    if action:
        qs = qs.filter(action_type__icontains=action)

    page_no = _safe_int(request.GET.get("page"), 1, 1, 10000)
    page_size = _safe_int(request.GET.get("page_size"), 50, 1, 200)
    page = Paginator(qs, page_size).get_page(page_no)

    items = [
        {
            "id": log.id,
            "operator": log.operator.username if log.operator else "-",
            "action_type": log.action_type,
            "object_type": log.object_type,
            "object_id": log.object_id,
            "detail": log.detail,
            "ip": log.ip,
            "created_at": log.created_at.strftime("%Y-%m-%d %H:%M:%S"),
        }
        for log in page.object_list
    ]
    return JsonResponse({
        "code": 200, "message": "ok",
        "data": {"items": items, "total": page.paginator.count, "page": page_no},
    })


@require_GET
def api_export(request):
    """GET /api/v1/export/prescriptions 处方清单导出（Excel）。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "导出:执行"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "处方清单"
    ws.append(["处方号", "患者", "医师", "诊断", "辨证", "处方类型", "状态", "缴费状态", "创建时间"])

    qs = (
        Prescription.objects.filter(is_deleted=False)
        .select_related("patient", "doctor__user")
        .order_by("-created_at")[:1000]
    )
    for p in qs:
        ws.append([
            p.presc_no, p.patient.name,
            p.doctor.user.get_full_name() or p.doctor.user.username,
            p.diagnosis, p.syndrome, p.get_presc_type_display(),
            p.get_status_display(), p.get_pay_status_display(),
            p.created_at.strftime("%Y-%m-%d %H:%M"),
        ])

    _log_export(request, "prescription", "导出处方清单")
    return _xlsx_response(wb, "prescriptions.xlsx")


@require_GET
def api_export_herbs(request):
    """GET /api/v1/export/herbs 药材台账导出（Excel）。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "导出:执行"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "药材台账"
    ws.append(["药材名称", "拼音首字母", "类别", "性味归经", "功效主治", "道地产地", "炮制规格",
               "剂量上限(g)", "常用剂量(g)", "毒性", "库存(kg)", "预警线(kg)"])

    for h in Herb.objects.all().order_by("pinyin_initial", "name"):
        ws.append([
            h.name, h.pinyin_initial, h.category, h.property, h.efficacy, h.origin, h.processing,
            float(h.max_dose) if h.max_dose else "", float(h.default_dose) if h.default_dose else "",
            "是" if h.is_toxic else "否", float(h.stock), float(h.stock_warning),
        ])

    _log_export(request, "herb", "导出药材台账")
    return _xlsx_response(wb, "herbs.xlsx")


@require_GET
def api_export_logs(request):
    """GET /api/v1/export/logs 操作日志导出（Excel）。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "导出:执行"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    from openpyxl import Workbook

    from .models import OperationLog

    wb = Workbook()
    ws = wb.active
    ws.title = "操作日志"
    ws.append(["操作人", "动作", "对象类型", "对象ID", "详情", "IP", "结果", "时间"])

    for log in OperationLog.objects.select_related("operator").order_by("-created_at")[:5000]:
        ws.append([
            log.operator.username if log.operator else "-", log.action_type, log.object_type,
            log.object_id, log.detail, log.ip, log.result,
            log.created_at.strftime("%Y-%m-%d %H:%M:%S"),
        ])

    _log_export(request, "operation_log", "导出操作日志")
    return _xlsx_response(wb, "operation_logs.xlsx")


@require_GET
def api_export_reviews(request):
    """GET /api/v1/export/reviews 审核记录导出（Excel）。"""
    if not request.user.is_authenticated:
        return JsonResponse({"code": 401, "message": "未登录"}, status=401)
    if not has_permission(request.user, "导出:执行"):
        return JsonResponse({"code": 403, "message": "无权限"}, status=403)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "审核记录"
    ws.append(["处方号", "审核级别", "审核人", "审核意见", "驳回原因", "不良反应备注", "审核时间"])

    rvs = ReviewRecord.objects.select_related("prescription", "pharmacist__user").order_by("-created_at")[:5000]
    for rv in rvs:
        ws.append([
            rv.prescription.presc_no,
            rv.get_review_level_display(),
            (rv.pharmacist.user.get_full_name() or rv.pharmacist.user.username) if rv.pharmacist else "-",
            rv.review_opinion, rv.reject_reason or "", rv.adverse_note or "",
            rv.created_at.strftime("%Y-%m-%d %H:%M:%S"),
        ])

    _log_export(request, "review_record", "导出审核记录")
    return _xlsx_response(wb, "review_records.xlsx")


# ============ 内部工具 ============

def _safe_int(raw, default, minimum, maximum):
    try:
        value = int(raw)
    except (TypeError, ValueError):
        return default
    return max(minimum, min(maximum, value))


def _log_export(request, object_type, detail):
    from .models import OperationLog

    OperationLog.objects.create(
        operator=request.user, action_type="EXPORT", object_type=object_type,
        object_id="all", detail=detail, ip=get_client_ip(request), result="success",
    )


def _xlsx_response(workbook, filename):
    import io

    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
